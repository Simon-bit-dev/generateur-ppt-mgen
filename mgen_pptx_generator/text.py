from __future__ import annotations

import re
from collections.abc import Iterable

from pptx.text.text import TextFrame

_BOLD_PATTERN = re.compile(r"\*\*(.+?)\*\*")


def _add_runs_with_optional_bold(paragraph, text: str) -> None:
    """Add text runs, converting **bold** markers into actual bold runs."""
    pos = 0
    for match in _BOLD_PATTERN.finditer(text):
        if match.start() > pos:
            paragraph.add_run().text = text[pos : match.start()]
        run = paragraph.add_run()
        run.text = match.group(1)
        run.font.bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run().text = text[pos:]
    if not text:
        paragraph.add_run().text = ""


def set_text_frame(
    text_frame: TextFrame,
    value: str | list[str],
    *,
    style: str = "plain",
) -> None:
    """Populate a text frame while preserving the layout's inherited styling.

    The MGEN templates use literal bullet and number prefixes rather than native
    PowerPoint bullets for most body layouts. This function follows that
    convention and never changes font sizes.
    """
    if isinstance(value, str):
        lines = value.splitlines() or [""]
    else:
        lines = [str(item) for item in value]

    if style == "bullet_text":
        lines = [line if line.startswith("• ") else f"• {line}" for line in lines]
    elif style == "numbered_text":
        lines = [
            line if re.match(r"^\d+[.)]\s", line) else f"{index}. {line}"
            for index, line in enumerate(lines, start=1)
        ]
    elif style != "plain":
        raise ValueError(f"Unknown text style: {style}")

    text_frame.clear()
    for index, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        _add_runs_with_optional_bold(paragraph, line)
