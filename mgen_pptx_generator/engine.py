from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Any

from pptx import Presentation

from .errors import GenerationError
from .models import SlideSpec
from .text import set_text_frame


def _layout_by_name(prs: Presentation, name: str):
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    raise GenerationError(f"Disposition absente de la template: {name}")


def _placeholder_by_idx(slide, idx: int):
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == idx:
            return shape
    raise GenerationError(f"Placeholder PH{idx} absent sur la disposition {slide.slide_layout.name}.")


def _clone_slide_number(layout, slide) -> None:
    """Clone the inherited slide-number field omitted by python-pptx add_slide()."""
    if any(shape.is_placeholder and shape.placeholder_format.idx == 2 for shape in slide.shapes):
        return
    for shape in layout.shapes:
        if shape.is_placeholder and shape.placeholder_format.idx == 2:
            slide.shapes._spTree.insert_element_before(deepcopy(shape.element), "p:extLst")
            return


def create_presentation(template_path: Path, specs: list[SlideSpec], output_path: Path) -> Presentation:
    prs = Presentation(str(template_path))
    if len(prs.slides) != 0:
        raise GenerationError("La template contient déjà des slides ordinaires.")

    for spec in specs:
        layout = _layout_by_name(prs, spec.layout_name)
        slide = prs.slides.add_slide(layout)
        if spec.clone_slide_number:
            _clone_slide_number(layout, slide)
        for placeholder in spec.placeholders:
            if placeholder.value in ("", []) and not placeholder.required:
                continue
            shape = _placeholder_by_idx(slide, placeholder.idx)
            if not shape.has_text_frame:
                raise GenerationError(f"PH{placeholder.idx} de {spec.layout_name} n’est pas textuel.")
            set_text_frame(shape.text_frame, placeholder.value, style=placeholder.style)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return prs


def inspect_generated(prs: Presentation, specs: list[SlideSpec]) -> dict[str, Any]:
    actual_layouts = [slide.slide_layout.name for slide in prs.slides]
    expected_layouts = [spec.layout_name for spec in specs]
    return {
        "slide_count": len(prs.slides),
        "expected_slide_count": len(specs),
        "layout_sequence_matches": actual_layouts == expected_layouts,
        "layout_sequence": actual_layouts,
        "roles": [spec.role for spec in specs],
    }
