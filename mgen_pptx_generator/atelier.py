from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any

from pptx import Presentation

from .errors import ValidationError
from .models import PlaceholderValue as PH
from .models import SlideSpec

ATELIER_ROOTS = {"VARIABLES_GLOBALES", "SLIDES", "ANOMALIES_TECHNIQUES"}
SCHEMA_ROOTS = {"VARIABLES_GLOBALES", "SLIDES"}
_VARIABLE_PATTERN = re.compile(r"\{\{([A-Z0-9_]+)\}\}")
_UNFILLED_MARKERS = ("__", "À RENSEIGNER", "À RENSEIGNER OU LAISSER VIDE")

EXPECTED_LAYOUTS = [
    "couverture",
    "Bienvenue!",
    "intervenants",
    "Accueil et objectifs",
    "sommaire",
    "liste objectifs",
    "brise glace",
    "transition",
    "apports théoriques",
    "VraiFaux transition",
    "VraiFaux?",
    "Faux !",
    "Vrai!",
    "exercices pratiques",
    "synthèse",
    "pour aller plus loin",
    "Vivoptim",
    "4e de couverture",
]

EXPECTED_PLACEHOLDERS: dict[str, set[int]] = {
    "couverture": {23, 24, 25, 26},
    "Bienvenue!": {21, 22},
    "intervenants": {2, 21, 25, 26},
    "Accueil et objectifs": {24, 25},
    "sommaire": {2, 25, 26},
    "liste objectifs": {2, 24, 25, 26},
    "brise glace": {2, 24, 25, 26, 27},
    "transition": {21, 23, 24},
    "apports théoriques": {2, 25, 26, 27, 28},
    "VraiFaux transition": {21},
    "VraiFaux?": {2, 21, 25, 26},
    "Faux !": {2, 25, 26, 27, 28},
    "Vrai!": {2, 21, 25, 26, 27},
    "exercices pratiques": {2, 25, 26, 27, 28},
    "synthèse": {2, 25, 26, 27},
    "pour aller plus loin": {2, 25, 26, 27},
    "Vivoptim": {2, 25, 26},
    "4e de couverture": {2, 26},
}

CORRECTION_RULES = {
    "MGEN-12": ("Faux !", {"PH25", "PH26", "PH27", "PH28"}),
    "MGEN-13": ("Vrai!", {"PH21", "PH25", "PH26", "PH27"}),
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise ValidationError(message)


def _required_field(
    mapping: Any,
    key: str,
    path: str,
    *,
    document: str = "YAML Atelier",
) -> Any:
    if not isinstance(mapping, dict):
        parent = path.rsplit(".", 1)[0] if "." in path else "la racine"
        raise ValidationError(
            f"Structure invalide dans le {document} : {parent} doit être un objet."
        )
    if key not in mapping:
        raise ValidationError(
            f"Champ obligatoire manquant dans le {document} : {path}."
        )
    return mapping[key]


def _required_mapping(
    mapping: Any,
    key: str,
    path: str,
    *,
    document: str = "YAML Atelier",
) -> dict[str, Any]:
    value = _required_field(mapping, key, path, document=document)
    _require(
        isinstance(value, dict),
        f"Le champ {path} du {document} doit être un objet.",
    )
    return value


def _required_list(
    mapping: Any,
    key: str,
    path: str,
    *,
    document: str = "YAML Atelier",
) -> list[Any]:
    value = _required_field(mapping, key, path, document=document)
    _require(
        isinstance(value, list),
        f"Le champ {path} du {document} doit être une liste.",
    )
    return value


def _contains_unfilled(value: Any) -> bool:
    if isinstance(value, str):
        return any(marker in value for marker in _UNFILLED_MARKERS)
    if isinstance(value, list):
        return any(_contains_unfilled(item) for item in value)
    if isinstance(value, dict):
        return any(_contains_unfilled(item) for item in value.values())
    return False


def _parse_ph_key(key: str) -> int:
    match = re.fullmatch(r"PH(\d+)", key)
    if not match:
        raise ValidationError(f"Clé de placeholder invalide: {key!r}.")
    return int(match.group(1))


def _resolve_variables(value: str, variables: dict[str, str]) -> str:
    def replace(match: re.Match[str]) -> str:
        name = match.group(1)
        if name not in variables:
            raise ValidationError(f"Variable globale inconnue: {name}.")
        return variables[name]

    resolved = _VARIABLE_PATTERN.sub(replace, value)
    if "{{" in resolved or "}}" in resolved:
        raise ValidationError(f"Variable globale mal formée ou non résolue dans: {value!r}.")
    return resolved


def _normalize_roots(data: dict[str, Any]) -> set[str]:
    return set(data.keys())


def validate_atelier_yaml(data: dict[str, Any], schema: dict[str, Any]) -> dict[str, int]:
    missing_roots = sorted(ATELIER_ROOTS - _normalize_roots(data))
    if missing_roots:
        raise ValidationError(
            "Champ obligatoire manquant dans le YAML Atelier : "
            f"{missing_roots[0]}."
        )
    _require(_normalize_roots(data) == ATELIER_ROOTS, "Le YAML Atelier doit conserver exactement les trois clés racines attendues.")

    missing_schema_roots = sorted(SCHEMA_ROOTS - set(schema))
    if missing_schema_roots:
        raise ValidationError(
            "Champ obligatoire manquant dans le schéma Atelier : "
            f"{missing_schema_roots[0]}."
        )
    _require(not _contains_unfilled(data), "Le YAML Atelier contient encore un marqueur de remplissage.")

    variables = _required_mapping(
        data,
        "VARIABLES_GLOBALES",
        "VARIABLES_GLOBALES",
    )
    for key in ("GLOBAL_FOOTER_LEFT", "GLOBAL_FOOTER_RIGHT"):
        _required_field(variables, key, f"VARIABLES_GLOBALES.{key}")
    _require(set(variables) == {"GLOBAL_FOOTER_LEFT", "GLOBAL_FOOTER_RIGHT"}, "Les deux variables globales Atelier doivent être conservées exactement.")
    _require(all(isinstance(value, str) and value.strip() for value in variables.values()), "Les variables globales Atelier doivent être des chaînes non vides.")

    anomalies = _required_mapping(
        data,
        "ANOMALIES_TECHNIQUES",
        "ANOMALIES_TECHNIQUES",
    )
    anomaly_status = _required_field(
        anomalies,
        "statut_global",
        "ANOMALIES_TECHNIQUES.statut_global",
    )
    declared_anomalies = _required_field(
        anomalies,
        "anomalies",
        "ANOMALIES_TECHNIQUES.anomalies",
    )
    _require(anomaly_status == "CONFORME", "Le YAML Atelier doit avoir un statut technique CONFORME.")
    _require(declared_anomalies == [], "Le YAML Atelier contient des anomalies techniques déclarées.")

    slides = _required_list(data, "SLIDES", "SLIDES")
    schema_slides = _required_list(
        schema,
        "SLIDES",
        "SLIDES",
        document="schéma Atelier",
    )
    _require(len(slides) == 33, "Le profil Atelier standard doit contenir exactement 33 slides.")
    _require(len(schema_slides) == 33, "Le schéma Atelier fourni doit décrire exactement 33 slides.")

    layout_counts: dict[str, int] = {}
    layout_sequence: list[str] = []
    for index, (slide, schema_slide) in enumerate(zip(slides, schema_slides), start=1):
        expected_id = f"S{index:02d}"
        slide_path = f"SLIDES[{index - 1}]"
        schema_slide_path = f"SLIDES[{index - 1}]"
        _require(
            isinstance(slide, dict),
            f"Le champ {slide_path} du YAML Atelier doit être un objet.",
        )
        _require(
            isinstance(schema_slide, dict),
            f"Le champ {schema_slide_path} du schéma Atelier doit être un objet.",
        )

        slide_id = _required_field(slide, "slide_id", f"{slide_path}.slide_id")
        number = _required_field(slide, "numero", f"{slide_path}.numero")
        part = _required_field(slide, "partie", f"{slide_path}.partie")
        status = _required_field(slide, "statut", f"{slide_path}.statut")
        code = _required_field(slide, "layout_code", f"{slide_path}.layout_code")
        name = _required_field(
            slide,
            "layout_name_exact",
            f"{slide_path}.layout_name_exact",
        )
        placeholders = _required_mapping(
            slide,
            "placeholders",
            f"{slide_path}.placeholders",
        )

        schema_id = _required_field(
            schema_slide,
            "slide_id",
            f"{schema_slide_path}.slide_id",
            document="schéma Atelier",
        )
        schema_number = _required_field(
            schema_slide,
            "numero",
            f"{schema_slide_path}.numero",
            document="schéma Atelier",
        )
        schema_part = _required_field(
            schema_slide,
            "partie",
            f"{schema_slide_path}.partie",
            document="schéma Atelier",
        )
        schema_code = _required_field(
            schema_slide,
            "layout_code",
            f"{schema_slide_path}.layout_code",
            document="schéma Atelier",
        )
        schema_name = _required_field(
            schema_slide,
            "layout_name_exact",
            f"{schema_slide_path}.layout_name_exact",
            document="schéma Atelier",
        )
        schema_placeholders = _required_mapping(
            schema_slide,
            "placeholders",
            f"{schema_slide_path}.placeholders",
            document="schéma Atelier",
        )

        _require(slide_id == expected_id, f"Slide {index}: slide_id attendu {expected_id}.")
        _require(number == index, f"Slide {expected_id}: numero attendu {index}.")
        _require(slide_id == schema_id, f"Slide {expected_id}: divergence avec le schéma.")
        _require(number == schema_number, f"Slide {expected_id}: numero divergent du schéma.")
        _require(part == schema_part, f"Slide {expected_id}: partie divergente du schéma.")
        _require(
            isinstance(status, str)
            and status in {"variable", "fixe — conserver intégralement"},
            f"Slide {expected_id}: statut invalide.",
        )

        if isinstance(schema_code, str) and schema_code.startswith("__MGEN_12"):
            _require(
                isinstance(code, str) and code in CORRECTION_RULES,
                f"Slide {expected_id}: correction attendue en MGEN-12 ou MGEN-13.",
            )
            expected_name, expected_keys = CORRECTION_RULES[code]
            _require(name == expected_name, f"Slide {expected_id}: nom de correction incompatible avec {code}.")
        else:
            _require(code == schema_code, f"Slide {expected_id}: layout_code divergent du schéma.")
            _require(name == schema_name, f"Slide {expected_id}: layout_name_exact divergent du schéma.")
            expected_keys = set(schema_placeholders)

        missing_placeholders = sorted(expected_keys - set(placeholders))
        if missing_placeholders:
            raise ValidationError(
                "Champ obligatoire manquant dans le YAML Atelier : "
                f"{slide_path}.placeholders.{missing_placeholders[0]}."
            )
        _require(
            set(placeholders) == expected_keys,
            f"Slide {expected_id}: clés de placeholders divergentes du schéma.",
        )

        _require(name in EXPECTED_LAYOUTS, f"Slide {expected_id}: disposition Atelier inconnue {name!r}.")
        allowed_idxs = EXPECTED_PLACEHOLDERS[name]
        for key, value in placeholders.items():
            _require(
                isinstance(key, str),
                f"Le chemin {slide_path}.placeholders contient une clé non textuelle.",
            )
            idx = _parse_ph_key(key)
            _require(idx in allowed_idxs, f"Slide {expected_id}: {key} n’existe pas sur la disposition {name}.")
            _require(isinstance(value, str), f"Le champ {slide_path}.placeholders.{key} doit contenir une chaîne.")
            _resolve_variables(value, variables)

        layout_counts[name] = layout_counts.get(name, 0) + 1
        layout_sequence.append(name)

    _require(layout_counts.get("apports théoriques") == 8, "L’atelier standard doit contenir exactement 8 slides d’apports théoriques.")
    _require(layout_counts.get("VraiFaux?") == 3, "L’atelier standard doit contenir exactement 3 questions Vrai ou Faux.")
    _require(layout_counts.get("Faux !", 0) + layout_counts.get("Vrai!", 0) == 3, "L’atelier standard doit contenir exactement 3 corrections.")
    _require(layout_counts.get("exercices pratiques") == 4, "L’atelier standard doit contenir exactement 4 exercices pratiques.")
    _require(layout_sequence[-3] == "pour aller plus loin", "La slide Pour aller plus loin doit précéder Vivoptim.")
    _require(layout_sequence[-2] == "Vivoptim", "Vivoptim doit être l’avant-dernière séquence institutionnelle.")
    _require(layout_sequence[-1] == "4e de couverture", "La quatrième de couverture doit être la dernière slide.")

    return {
        "total_slides": 33,
        "theoretical_slides": 8,
        "quiz_questions": 3,
        "quiz_corrections": 3,
        "exercise_slides": 4,
    }


def audit_atelier_template(path: Path, data: dict[str, Any], schema: dict[str, Any]) -> dict[str, Any]:
    del data, schema  # The YAML family has no embedded template fingerprint in this version.
    actual_sha = sha256_file(path)
    prs = Presentation(str(path))
    _require(len(prs.slide_masters) == 1, "La template Atelier doit contenir exactement un masque.")
    _require(len(prs.slide_layouts) == 18, "La template Atelier doit contenir exactement 18 dispositions.")
    _require(len(prs.slides) == 0, "La template Atelier doit être dépourvue de slides ordinaires.")

    names = [layout.name for layout in prs.slide_layouts]
    _require(names == EXPECTED_LAYOUTS, "Les noms ou l’ordre des 18 dispositions Atelier diffèrent du registre attendu.")
    _require(len(set(names)) == len(names), "Les noms de dispositions Atelier doivent être uniques.")

    for layout in prs.slide_layouts:
        actual = {shape.placeholder_format.idx for shape in layout.shapes if shape.is_placeholder}
        expected = EXPECTED_PLACEHOLDERS[layout.name]
        _require(actual == expected, f"Disposition {layout.name}: placeholders attendus {sorted(expected)}, obtenus {sorted(actual)}.")

    objectives_layout = prs.slide_layouts[names.index("liste objectifs")]
    fixed_texts = [shape.text for shape in objectives_layout.shapes if hasattr(shape, "text") and shape.text]
    _require("Les objectifs de l’atelier" in fixed_texts, "Le titre fixe de la disposition Objectifs Atelier est non conforme.")

    return {
        "sha256": actual_sha,
        "slide_master_count": len(prs.slide_masters),
        "layout_count": len(prs.slide_layouts),
        "ordinary_slide_count": len(prs.slides),
        "layout_names": names,
    }


def build_atelier_specs(data: dict[str, Any]) -> list[SlideSpec]:
    variables = data["VARIABLES_GLOBALES"]
    specs: list[SlideSpec] = []
    for slide in data["SLIDES"]:
        placeholders: list[PH] = []
        for key, raw_value in slide["placeholders"].items():
            idx = _parse_ph_key(key)
            value = _resolve_variables(raw_value, variables)
            placeholders.append(PH(idx, value, style="plain", required=True))
        specs.append(
            SlideSpec(
                slide["layout_name_exact"],
                f"{slide['slide_id']} — {slide['partie']}",
                tuple(placeholders),
                clone_slide_number=True,
            )
        )
    return specs
