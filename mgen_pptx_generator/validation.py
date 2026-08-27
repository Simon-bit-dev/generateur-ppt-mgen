from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

import yaml
from pptx import Presentation

from .errors import ValidationError

EXPECTED_LAYOUTS = [
    "01_Couverture",
    "02_Bienvenue",
    "03_Intervenants",
    "04_Sommaire",
    "05_Transition_Accueil",
    "06_Objectifs",
    "07_Brise_glace",
    "08_Transition_Repères_théoriques",
    "09_Corps_Repères_théoriques",
    "10_Transition_Conseils_pratiques",
    "11_Corps_Conseils_pratiques",
    "12_Transition_Testez_connaissances",
    "13_Vrai_ou_Faux",
    "14_Faux",
    "15_Vrai",
    "16_Transition_Passer_action",
    "17_Corps_Passer_action",
    "18_Transition_Conclusion",
    "19_Synthèse",
    "20_Pour_aller_plus_loin",
    "21_Vivoptim",
    "22_4ème_couverture",
]

EXPECTED_PLACEHOLDERS: dict[str, set[int]] = {
    "01_Couverture": {23, 24, 25, 26},
    "02_Bienvenue": {21, 22},
    "03_Intervenants": {2, 21, 25, 26},
    "04_Sommaire": {2, 25, 26},
    "05_Transition_Accueil": {24},
    "06_Objectifs": {2, 24, 25, 26},
    "07_Brise_glace": {2, 24, 25, 27, 28},
    "08_Transition_Repères_théoriques": {24},
    "09_Corps_Repères_théoriques": {2, 25, 26, 28, 29},
    "10_Transition_Conseils_pratiques": {24},
    "11_Corps_Conseils_pratiques": {2, 25, 26, 28, 29},
    "12_Transition_Testez_connaissances": set(),
    "13_Vrai_ou_Faux": {2, 21, 25, 26},
    "14_Faux": {2, 25, 26, 28, 29},
    "15_Vrai": {2, 21, 25, 26, 27},
    "16_Transition_Passer_action": {24},
    "17_Corps_Passer_action": {2, 25, 26, 28, 29},
    "18_Transition_Conclusion": {24},
    "19_Synthèse": {2, 25, 26, 27},
    "20_Pour_aller_plus_loin": {2, 25, 26, 27},
    "21_Vivoptim": {2, 25, 26},
    "22_4ème_couverture": {2, 26},
}

REQUIRED_ROOTS = {"schema_version", "template", "presentation", "opening", "sections", "validation_contract"}
MARKERS = ("À RENSEIGNER", "À RENSEIGNER OU LAISSER VIDE")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_yaml(path: Path) -> dict[str, Any]:
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, yaml.YAMLError) as exc:
        raise ValidationError(f"YAML illisible ou invalide: {exc}") from exc
    if not isinstance(data, dict):
        raise ValidationError("La racine du YAML doit être un objet.")
    return data


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise ValidationError(message)


def _required_field(mapping: Any, key: str, path: str) -> Any:
    if not isinstance(mapping, dict):
        parent = path.rsplit(".", 1)[0] if "." in path else "la racine"
        raise ValidationError(
            f"Structure invalide dans le YAML Conférence : {parent} doit être un objet."
        )
    if key not in mapping:
        raise ValidationError(
            f"Champ obligatoire manquant dans le YAML Conférence : {path}."
        )
    return mapping[key]


def _required_mapping(mapping: Any, key: str, path: str) -> dict[str, Any]:
    value = _required_field(mapping, key, path)
    _require(
        isinstance(value, dict),
        f"Le champ {path} du YAML Conférence doit être un objet.",
    )
    return value


def _required_list(mapping: Any, key: str, path: str) -> list[Any]:
    value = _required_field(mapping, key, path)
    _require(
        isinstance(value, list),
        f"Le champ {path} du YAML Conférence doit être une liste.",
    )
    return value


def _required_int(mapping: Any, key: str, path: str) -> int:
    value = _required_field(mapping, key, path)
    _require(
        isinstance(value, int) and not isinstance(value, bool),
        f"Le champ {path} du YAML Conférence doit être un entier.",
    )
    return value


def _contains_marker(value: Any) -> bool:
    if isinstance(value, str):
        return any(marker in value for marker in MARKERS)
    if isinstance(value, list):
        return any(_contains_marker(item) for item in value)
    if isinstance(value, dict):
        return any(_contains_marker(item) for item in value.values())
    return False


def validate_semantic_yaml(data: dict[str, Any]) -> dict[str, int]:
    missing_roots = sorted(REQUIRED_ROOTS - set(data))
    if missing_roots:
        raise ValidationError(
            "Champ obligatoire manquant dans le YAML Conférence : "
            f"{missing_roots[0]}."
        )
    _require(
        set(data) == REQUIRED_ROOTS,
        "Les six clés racines du schéma doivent être conservées exactement.",
    )

    schema_version = _required_field(data, "schema_version", "schema_version")
    template = _required_mapping(data, "template", "template")
    presentation = _required_mapping(data, "presentation", "presentation")
    opening = _required_mapping(data, "opening", "opening")
    sections = _required_mapping(data, "sections", "sections")
    contract = _required_mapping(
        data,
        "validation_contract",
        "validation_contract",
    )

    _require(schema_version == "1.0", "schema_version doit valoir 1.0.")
    _require(not _contains_marker(data), "Le YAML contient encore un marqueur À RENSEIGNER.")

    template_sha = _required_field(
        template,
        "template_sha256",
        "template.template_sha256",
    )
    _require(
        isinstance(template_sha, str) and bool(template_sha.strip()),
        "Le champ template.template_sha256 doit être une chaîne non vide.",
    )
    _required_int(
        template,
        "expected_slide_master_count",
        "template.expected_slide_master_count",
    )
    _required_int(
        template,
        "expected_layout_count",
        "template.expected_layout_count",
    )

    architecture = _required_field(
        presentation,
        "architecture_variant",
        "presentation.architecture_variant",
    )
    _require(
        isinstance(architecture, str) and bool(architecture.strip()),
        "Le champ presentation.architecture_variant doit être une chaîne non vide.",
    )
    allowed = _required_mapping(
        contract,
        "allowed_architectures",
        "validation_contract.allowed_architectures",
    )
    _require(architecture in allowed, f"Architecture non autorisée: {architecture!r}.")
    counts = _required_mapping(
        allowed,
        architecture,
        f"validation_contract.allowed_architectures.{architecture}",
    )
    for key in (
        "theoretical_slides",
        "practical_slides",
        "action_slides",
        "total_slides",
    ):
        _required_int(
            counts,
            key,
            f"validation_contract.allowed_architectures.{architecture}.{key}",
        )

    allowed_formats = _required_list(
        contract,
        "allowed_formats",
        "validation_contract.allowed_formats",
    )
    limits = _required_mapping(
        contract,
        "limits",
        "validation_contract.limits",
    )
    verdict_mapping = _required_mapping(
        contract,
        "verdict_mapping",
        "validation_contract.verdict_mapping",
    )
    limit_values = {
        key: _required_int(limits, key, f"validation_contract.limits.{key}")
        for key in (
            "theoretical_body_items_max",
            "practical_body_items_max",
            "quiz_questions_exact",
            "quiz_explanation_items_max",
            "action_body_items_max",
            "synthesis_items_min",
            "synthesis_items_max",
            "further_resources_max",
        )
    }

    support_format = _required_field(presentation, "format", "presentation.format")
    title = _required_field(presentation, "title", "presentation.title")
    subtitle = _required_field(presentation, "subtitle", "presentation.subtitle")
    date = _required_field(presentation, "date", "presentation.date")
    _require(support_format in allowed_formats, "Format de support non autorisé.")
    _require(isinstance(title, str) and title.strip(), "Le titre est obligatoire.")
    _require(isinstance(subtitle, str), "Le sous-titre doit être une chaîne, éventuellement vide.")
    _require(isinstance(date, str) and date.strip(), "La date est obligatoire.")

    speaker_identification = _required_field(
        opening,
        "speaker_identification",
        "opening.speaker_identification",
    )
    welcome_sentence = _required_field(
        opening,
        "welcome_sentence",
        "opening.welcome_sentence",
    )
    _require(speaker_identification == "À compléter", "speaker_identification doit valoir exactement À compléter.")
    _require(
        isinstance(welcome_sentence, str) and bool(welcome_sentence.strip()),
        "La phrase de bienvenue est obligatoire.",
    )

    accueil = _required_mapping(
        sections,
        "accueil_objectifs",
        "sections.accueil_objectifs",
    )
    accueil_transition = _required_field(
        accueil,
        "transition_sentence",
        "sections.accueil_objectifs.transition_sentence",
    )
    objectives = _required_list(
        accueil,
        "objectives",
        "sections.accueil_objectifs.objectives",
    )
    _require(3 <= len(objectives) <= 4, "Il faut entre 3 et 4 objectifs.")
    _require(all(isinstance(x, str) and x.strip() for x in objectives), "Chaque objectif doit être non vide.")
    ice = _required_mapping(
        accueil,
        "icebreaker",
        "sections.accueil_objectifs.icebreaker",
    )
    ice_question = _required_field(
        ice,
        "question",
        "sections.accueil_objectifs.icebreaker.question",
    )
    ice_instruction = _required_field(
        ice,
        "instruction",
        "sections.accueil_objectifs.icebreaker.instruction",
    )
    _require(
        isinstance(ice_question, str)
        and bool(ice_question.strip())
        and isinstance(ice_instruction, str)
        and bool(ice_instruction.strip()),
        "Le brise-glace doit comporter une question et une consigne.",
    )

    section_transitions: dict[str, Any] = {}

    def validate_slots(name: str, expected_enabled: int, min_items: int, max_items: int) -> None:
        section_path = f"sections.{name}"
        section = _required_mapping(sections, name, section_path)
        section_transitions[name] = _required_field(
            section,
            "transition_sentence",
            f"{section_path}.transition_sentence",
        )
        slots = _required_list(section, "slides", f"{section_path}.slides")
        enabled_count = 0
        for expected_slot, slot in enumerate(slots, start=1):
            slot_path = f"{section_path}.slides[{expected_slot - 1}]"
            _require(
                isinstance(slot, dict),
                f"Le champ {slot_path} du YAML Conférence doit être un objet.",
            )
            slot_number = _required_field(slot, "slot", f"{slot_path}.slot")
            is_enabled = _required_field(slot, "enabled", f"{slot_path}.enabled")
            title = _required_field(slot, "title", f"{slot_path}.title")
            body = _required_field(slot, "body", f"{slot_path}.body")
            _require(slot_number == expected_slot, f"{name}: slots non continus ou modifiés.")
            _require(
                isinstance(is_enabled, bool),
                f"Le champ {slot_path}.enabled doit être un booléen.",
            )
            if is_enabled:
                enabled_count += 1
                _require(isinstance(title, str) and title.strip(), f"{name} slot {expected_slot}: titre obligatoire.")
                _require(isinstance(body, list) and min_items <= len(body) <= max_items, f"{name} slot {expected_slot}: nombre d’éléments invalide.")
                _require(all(isinstance(x, str) and x.strip() for x in body), f"{name} slot {expected_slot}: éléments vides interdits.")
            else:
                _require(title == "" and body == [], f"{name} slot {expected_slot}: un slot désactivé doit être vide.")
        _require(
            enabled_count == expected_enabled,
            f"{name}: nombre de slots activés incohérent avec l’architecture.",
        )

    validate_slots("reperes_theoriques", counts["theoretical_slides"], 3, limit_values["theoretical_body_items_max"])
    validate_slots("conseils_pratiques", counts["practical_slides"], 3, limit_values["practical_body_items_max"])
    validate_slots("passer_action", counts["action_slides"], 3, limit_values["action_body_items_max"])

    quiz = _required_mapping(
        sections,
        "testez_connaissances",
        "sections.testez_connaissances",
    )
    questions = _required_list(
        quiz,
        "questions",
        "sections.testez_connaissances.questions",
    )
    _require(len(questions) == limit_values["quiz_questions_exact"], "Il faut exactement trois questions Vrai ou Faux.")
    for expected_number, question in enumerate(questions, start=1):
        question_path = f"sections.testez_connaissances.questions[{expected_number - 1}]"
        _require(
            isinstance(question, dict),
            f"Le champ {question_path} du YAML Conférence doit être un objet.",
        )
        number = _required_field(question, "number", f"{question_path}.number")
        verdict = _required_field(question, "verdict", f"{question_path}.verdict")
        statement = _required_field(question, "statement", f"{question_path}.statement")
        correction = _required_field(
            question,
            "correction_message",
            f"{question_path}.correction_message",
        )
        explanation = _required_list(
            question,
            "explanation",
            f"{question_path}.explanation",
        )
        _require(number == expected_number, "Les numéros des questions doivent rester continus.")
        _require(isinstance(verdict, str) and verdict in verdict_mapping, f"Question {expected_number}: verdict invalide.")
        _require(isinstance(statement, str) and bool(statement.strip()), f"Question {expected_number}: affirmation vide.")
        _require(isinstance(correction, str) and bool(correction.strip()), f"Question {expected_number}: correction vide.")
        _require(2 <= len(explanation) <= limit_values["quiz_explanation_items_max"], f"Question {expected_number}: explication invalide.")
        _require(
            all(isinstance(item, str) and item.strip() for item in explanation),
            f"Question {expected_number}: les éléments d’explication doivent être non vides.",
        )

    conclusion = _required_mapping(sections, "conclusion", "sections.conclusion")
    conclusion_transition = _required_field(
        conclusion,
        "transition_sentence",
        "sections.conclusion.transition_sentence",
    )
    synthesis = _required_list(
        conclusion,
        "synthesis",
        "sections.conclusion.synthesis",
    )
    resources = _required_list(
        conclusion,
        "further_resources",
        "sections.conclusion.further_resources",
    )
    _require(limit_values["synthesis_items_min"] <= len(synthesis) <= limit_values["synthesis_items_max"], "La synthèse doit comporter 4 ou 5 éléments.")
    _require(2 <= len(resources) <= limit_values["further_resources_max"], "Pour aller plus loin doit comporter 2 à 4 ressources.")
    _require(
        all(isinstance(item, str) and item.strip() for item in synthesis),
        "Chaque élément de synthèse doit être une chaîne non vide.",
    )
    _require(
        all(isinstance(item, str) and item.strip() for item in resources),
        "Chaque ressource doit être une chaîne non vide.",
    )

    transition_paths = [
        accueil_transition,
        section_transitions["reperes_theoriques"],
        section_transitions["conseils_pratiques"],
        section_transitions["passer_action"],
        conclusion_transition,
    ]
    _require(all(isinstance(x, str) and x.strip() for x in transition_paths), "Toutes les transitions variables doivent être renseignées.")
    return {k: int(v) for k, v in counts.items()}


def audit_template(path: Path, data: dict[str, Any]) -> dict[str, Any]:
    actual_sha = sha256_file(path)
    template = _required_mapping(data, "template", "template")
    expected_sha = _required_field(
        template,
        "template_sha256",
        "template.template_sha256",
    )
    expected_master_count = _required_int(
        template,
        "expected_slide_master_count",
        "template.expected_slide_master_count",
    )
    expected_layout_count = _required_int(
        template,
        "expected_layout_count",
        "template.expected_layout_count",
    )
    _require(actual_sha == expected_sha, f"Empreinte de template différente: attendu {expected_sha}, obtenu {actual_sha}.")

    prs = Presentation(str(path))
    _require(len(prs.slide_masters) == expected_master_count, "Nombre de masques incorrect.")
    _require(len(prs.slide_layouts) == expected_layout_count, "Nombre de dispositions incorrect.")
    _require(len(prs.slides) == 0, "La template doit être dépourvue de slides ordinaires.")

    names = [layout.name for layout in prs.slide_layouts]
    _require(names == EXPECTED_LAYOUTS, "Les noms ou l’ordre des 22 dispositions diffèrent du registre attendu.")
    _require(len(set(names)) == len(names), "Les noms de dispositions doivent être uniques.")

    for layout in prs.slide_layouts:
        actual = {shape.placeholder_format.idx for shape in layout.shapes if shape.is_placeholder}
        expected = EXPECTED_PLACEHOLDERS[layout.name]
        _require(actual == expected, f"Disposition {layout.name}: placeholders attendus {sorted(expected)}, obtenus {sorted(actual)}.")

    objectives_layout = prs.slide_layouts[names.index("06_Objectifs")]
    fixed_texts = [shape.text for shape in objectives_layout.shapes if hasattr(shape, "text") and shape.text]
    _require("Les objectifs de la conférence" in fixed_texts, "Le titre fixe de la disposition Objectifs est non conforme.")

    return {
        "sha256": actual_sha,
        "slide_master_count": len(prs.slide_masters),
        "layout_count": len(prs.slide_layouts),
        "ordinary_slide_count": len(prs.slides),
        "layout_names": names,
    }
