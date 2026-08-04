from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

import yaml
from pptx import Presentation

from .errors import ValidationError

EXPECTED_LAYOUTS = [
    "01_Couverture",
    "02_Bienvenue",
    "03_Intervenants",
    "04_Sommaire",
    "05_Transition_Accueil",
    "06_Objectifs",
    "07_Brise_glace",
    "08_Transition_Repères_théoriques",
    "09_Corps_Repères_théoriques",
    "10_Transition_Conseils_pratiques",
    "11_Corps_Conseils_pratiques",
    "12_Transition_Testez_connaissances",
    "13_Vrai_ou_Faux",
    "14_Faux",
    "15_Vrai",
    "16_Transition_Passer_action",
    "17_Corps_Passer_action",
    "18_Transition_Conclusion",
    "19_Synthèse",
    "20_Pour_aller_plus_loin",
    "21_Vivoptim",
    "22_4ème_couverture",
]

EXPECTED_PLACEHOLDERS: dict[str, set[int]] = {
    "01_Couverture": {23, 24, 25, 26},
    "02_Bienvenue": {21, 22},
    "03_Intervenants": {2, 21, 25, 26},
    "04_Sommaire": {2, 25, 26},
    "05_Transition_Accueil": {24},
    "06_Objectifs": {2, 24, 25, 26},
    "07_Brise_glace": {2, 24, 25, 27, 28},
    "08_Transition_Repères_théoriques": {24},
    "09_Corps_Repères_théoriques": {2, 25, 26, 28, 29},
    "10_Transition_Conseils_pratiques": {24},
    "11_Corps_Conseils_pratiques": {2, 25, 26, 28, 29},
    "12_Transition_Testez_connaissances": set(),
    "13_Vrai_ou_Faux": {2, 21, 25, 26},
    "14_Faux": {2, 25, 26, 28, 29},
    "15_Vrai": {2, 21, 25, 26, 27},
    "16_Transition_Passer_action": {24},
    "17_Corps_Passer_action": {2, 25, 26, 28, 29},
    "18_Transition_Conclusion": {24},
    "19_Synthèse": {2, 25, 26, 27},
    "20_Pour_aller_plus_loin": {2, 25, 26, 27},
    "21_Vivoptim": {2, 25, 26},
    "22_4ème_couverture": {2, 26},
}

REQUIRED_ROOTS = {"schema_version", "template", "presentation", "opening", "sections", "validation_contract"}
MARKERS = ("À RENSEIGNER", "À RENSEIGNER OU LAISSER VIDE")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_yaml(path: Path) -> dict[str, Any]:
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
    except Exception as exc:  # pragma: no cover - defensive
        raise ValidationError(f"YAML illisible ou invalide: {exc}") from exc
    if not isinstance(data, dict):
        raise ValidationError("La racine du YAML doit être un objet.")
    return data


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise ValidationError(message)


def _contains_marker(value: Any) -> bool:
    if isinstance(value, str):
        return any(marker in value for marker in MARKERS)
    if isinstance(value, list):
        return any(_contains_marker(item) for item in value)
    if isinstance(value, dict):
        return any(_contains_marker(item) for item in value.values())
    return False


def validate_semantic_yaml(data: dict[str, Any]) -> dict[str, int]:
    _require(set(data.keys()) == REQUIRED_ROOTS, "Les six clés racines du schéma doivent être conservées exactement.")
    _require(data.get("schema_version") == "1.0", "schema_version doit valoir 1.0.")
    _require(not _contains_marker(data), "Le YAML contient encore un marqueur À RENSEIGNER.")

    presentation = data["presentation"]
    contract = data["validation_contract"]
    architecture = presentation.get("architecture_variant")
    allowed = contract["allowed_architectures"]
    _require(architecture in allowed, f"Architecture non autorisée: {architecture!r}.")
    counts = allowed[architecture]

    _require(presentation.get("format") in contract["allowed_formats"], "Format de support non autorisé.")
    _require(isinstance(presentation.get("title"), str) and presentation["title"].strip(), "Le titre est obligatoire.")
    _require(isinstance(presentation.get("subtitle"), str), "Le sous-titre doit être une chaîne, éventuellement vide.")
    _require(isinstance(presentation.get("date"), str) and presentation["date"].strip(), "La date est obligatoire.")
    _require(data["opening"].get("speaker_identification") == "À compléter", "speaker_identification doit valoir exactement À compléter.")
    _require(bool(data["opening"].get("welcome_sentence", "").strip()), "La phrase de bienvenue est obligatoire.")

    sections = data["sections"]
    objectives = sections["accueil_objectifs"]["objectives"]
    _require(isinstance(objectives, list) and 3 <= len(objectives) <= 4, "Il faut entre 3 et 4 objectifs.")
    _require(all(isinstance(x, str) and x.strip() for x in objectives), "Chaque objectif doit être non vide.")
    ice = sections["accueil_objectifs"]["icebreaker"]
    _require(bool(ice.get("question", "").strip()) and bool(ice.get("instruction", "").strip()), "Le brise-glace doit comporter une question et une consigne.")

    def validate_slots(name: str, expected_enabled: int, min_items: int, max_items: int) -> None:
        slots = sections[name]["slides"]
        enabled = [slot for slot in slots if slot.get("enabled") is True]
        _require(len(enabled) == expected_enabled, f"{name}: nombre de slots activés incohérent avec l’architecture.")
        for expected_slot, slot in enumerate(slots, start=1):
            _require(slot.get("slot") == expected_slot, f"{name}: slots non continus ou modifiés.")
            title = slot.get("title")
            body = slot.get("body")
            if slot.get("enabled") is True:
                _require(isinstance(title, str) and title.strip(), f"{name} slot {expected_slot}: titre obligatoire.")
                _require(isinstance(body, list) and min_items <= len(body) <= max_items, f"{name} slot {expected_slot}: nombre d’éléments invalide.")
                _require(all(isinstance(x, str) and x.strip() for x in body), f"{name} slot {expected_slot}: éléments vides interdits.")
            else:
                _require(title == "" and body == [], f"{name} slot {expected_slot}: un slot désactivé doit être vide.")

    validate_slots("reperes_theoriques", counts["theoretical_slides"], 3, contract["limits"]["theoretical_body_items_max"])
    validate_slots("conseils_pratiques", counts["practical_slides"], 3, contract["limits"]["practical_body_items_max"])
    validate_slots("passer_action", counts["action_slides"], 3, contract["limits"]["action_body_items_max"])

    questions = sections["testez_connaissances"]["questions"]
    _require(len(questions) == contract["limits"]["quiz_questions_exact"], "Il faut exactement trois questions Vrai ou Faux.")
    for expected_number, question in enumerate(questions, start=1):
        _require(question.get("number") == expected_number, "Les numéros des questions doivent rester continus.")
        _require(question.get("verdict") in contract["verdict_mapping"], f"Question {expected_number}: verdict invalide.")
        _require(bool(question.get("statement", "").strip()), f"Question {expected_number}: affirmation vide.")
        _require(bool(question.get("correction_message", "").strip()), f"Question {expected_number}: correction vide.")
        explanation = question.get("explanation")
        _require(isinstance(explanation, list) and 2 <= len(explanation) <= contract["limits"]["quiz_explanation_items_max"], f"Question {expected_number}: explication invalide.")

    conclusion = sections["conclusion"]
    synthesis = conclusion["synthesis"]
    resources = conclusion["further_resources"]
    _require(contract["limits"]["synthesis_items_min"] <= len(synthesis) <= contract["limits"]["synthesis_items_max"], "La synthèse doit comporter 4 ou 5 éléments.")
    _require(2 <= len(resources) <= contract["limits"]["further_resources_max"], "Pour aller plus loin doit comporter 2 à 4 ressources.")

    transition_paths = [
        sections["accueil_objectifs"]["transition_sentence"],
        sections["reperes_theoriques"]["transition_sentence"],
        sections["conseils_pratiques"]["transition_sentence"],
        sections["passer_action"]["transition_sentence"],
        sections["conclusion"]["transition_sentence"],
    ]
    _require(all(isinstance(x, str) and x.strip() for x in transition_paths), "Toutes les transitions variables doivent être renseignées.")
    return {k: int(v) for k, v in counts.items()}


def audit_template(path: Path, data: dict[str, Any]) -> dict[str, Any]:
    actual_sha = sha256_file(path)
    expected_sha = data["template"]["template_sha256"]
    _require(actual_sha == expected_sha, f"Empreinte de template différente: attendu {expected_sha}, obtenu {actual_sha}.")

    prs = Presentation(str(path))
    _require(len(prs.slide_masters) == data["template"]["expected_slide_master_count"], "Nombre de masques incorrect.")
    _require(len(prs.slide_layouts) == data["template"]["expected_layout_count"], "Nombre de dispositions incorrect.")
    _require(len(prs.slides) == 0, "La template doit être dépourvue de slides ordinaires.")

    names = [layout.name for layout in prs.slide_layouts]
    _require(names == EXPECTED_LAYOUTS, "Les noms ou l’ordre des 22 dispositions diffèrent du registre attendu.")
    _require(len(set(names)) == len(names), "Les noms de dispositions doivent être uniques.")

    for layout in prs.slide_layouts:
        actual = {shape.placeholder_format.idx for shape in layout.shapes if shape.is_placeholder}
        expected = EXPECTED_PLACEHOLDERS[layout.name]
        _require(actual == expected, f"Disposition {layout.name}: placeholders attendus {sorted(expected)}, obtenus {sorted(actual)}.")

    objectives_layout = prs.slide_layouts[names.index("06_Objectifs")]
    fixed_texts = [shape.text for shape in objectives_layout.shapes if hasattr(shape, "text") and shape.text]
    _require("Les objectifs de la conférence" in fixed_texts, "Le titre fixe de la disposition Objectifs est non conforme.")

    return {
        "sha256": actual_sha,
        "slide_master_count": len(prs.slide_masters),
        "layout_count": len(prs.slide_layouts),
        "ordinary_slide_count": len(prs.slides),
        "layout_names": names,
    }
