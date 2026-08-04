from __future__ import annotations

import argparse
import subprocess
import sys
from pathlib import Path

from pptx import Presentation

from .atelier import audit_atelier_template, build_atelier_specs, validate_atelier_yaml
from .conference import build_conference_specs
from .engine import create_presentation, inspect_generated
from .errors import MgenGeneratorError, ValidationError
from .profiles import detect_support_type
from .report import write_markdown_report
from .validation import audit_template as audit_conference_template
from .validation import load_yaml, validate_semantic_yaml


def _compare_reference(reference: Path, generated: Presentation) -> dict[str, object]:
    ref = Presentation(str(reference))
    ref_layouts = [slide.slide_layout.name for slide in ref.slides]
    gen_layouts = [slide.slide_layout.name for slide in generated.slides]

    def normalize(layout: str) -> str:
        correction_layouts = {"14_Faux", "15_Vrai", "Faux !", "Vrai!"}
        return "CORRECTION_VRAI_OU_FAUX" if layout in correction_layouts else layout

    return {
        "reference_slide_count": len(ref.slides),
        "same_slide_count": len(ref.slides) == len(generated.slides),
        "same_layout_sequence": [normalize(x) for x in ref_layouts] == [normalize(x) for x in gen_layouts],
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Génère un PowerPoint MGEN depuis un YAML Conférence ou Atelier validé.")
    parser.add_argument("--yaml", required=True, type=Path, help="YAML final Conférence ou Atelier")
    parser.add_argument("--template", required=True, type=Path, help="Template PowerPoint MGEN correspondante")
    parser.add_argument("--schema", type=Path, help="Schéma YAML Atelier obligatoire pour le profil Atelier")
    parser.add_argument("--output", required=True, type=Path, help="PowerPoint de sortie")
    parser.add_argument("--report", type=Path, help="Rapport Markdown de contrôle")
    parser.add_argument("--reference", type=Path, help="PowerPoint témoin facultatif")
    parser.add_argument("--render-dir", type=Path, help="Dossier de rendu PNG facultatif")
    parser.add_argument("--skip-render-checks", action="store_true", help="Désactive le rendu et le contrôle hors-canevas")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        data = load_yaml(args.yaml)
        try:
            support_type = detect_support_type(data)
        except ValueError as exc:
            raise ValidationError(str(exc)) from exc

        if support_type == "conference":
            counts = validate_semantic_yaml(data)
            template_audit = audit_conference_template(args.template, data)
            specs = build_conference_specs(data)
        else:
            if args.schema is None:
                raise ValidationError("Le profil Atelier exige l’argument --schema.")
            schema = load_yaml(args.schema)
            counts = validate_atelier_yaml(data, schema)
            template_audit = audit_atelier_template(args.template, data, schema)
            specs = build_atelier_specs(data)

        if len(specs) != counts["total_slides"]:
            raise MgenGeneratorError(f"Le moteur a construit {len(specs)} slides au lieu de {counts['total_slides']}.")
        prs = create_presentation(args.template, specs, args.output)
        generation = inspect_generated(prs, specs)
        reference = _compare_reference(args.reference, prs) if args.reference else None

        render_status = "NON EXÉCUTÉ"
        overflow_status = "NON EXÉCUTÉ"
        warnings: list[str] = []
        if not args.skip_render_checks:
            render_dir = args.render_dir or args.output.with_suffix("")
            render_script = Path("/home/oai/skills/slides/container_tools/render_slides.py")
            overflow_script = Path("/home/oai/skills/slides/container_tools/slides_test.py")
            render = subprocess.run(
                [sys.executable, str(render_script), str(args.output), "--output_dir", str(render_dir)],
                capture_output=True,
                text=True,
            )
            render_status = "OK" if render.returncode == 0 else "ÉCHEC"
            if render.returncode != 0:
                warnings.append(f"Rendu impossible: {render.stderr.strip() or render.stdout.strip()}")
            overflow = subprocess.run(
                [sys.executable, str(overflow_script), str(args.output)],
                capture_output=True,
                text=True,
            )
            overflow_status = "OK" if overflow.returncode == 0 else "ÉCHEC"
            if overflow.returncode != 0:
                warnings.append(f"Contrôle hors-canevas en échec: {overflow.stderr.strip() or overflow.stdout.strip()}")

        if args.report:
            write_markdown_report(
                args.report,
                support_type=support_type,
                yaml_path=args.yaml,
                template_path=args.template,
                output_path=args.output,
                template_audit=template_audit,
                validation_counts=counts,
                generation=generation,
                reference_comparison=reference,
                render_status=render_status,
                overflow_status=overflow_status,
                warnings=warnings,
            )
        print(f"Profil détecté: {support_type}")
        print(f"PowerPoint généré: {args.output}")
        if args.report:
            print(f"Rapport généré: {args.report}")
        return 0
    except MgenGeneratorError as exc:
        print(f"ERREUR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
